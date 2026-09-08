from __future__ import annotations

import csv
from collections.abc import Iterable
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from models import BenefitType, EmployeeRole, RdvSubmission
from utils import (
    calculate_rdv_totals,
    format_brl,
    format_date,
    format_long_date,
    protocol,
)


def _entry_rows(rdv: RdvSubmission) -> list[list[object]]:
    return [
        [
            protocol(rdv.id),
            rdv.employee.name,
            rdv.employee.role.value,
            format_date(entry.date),
            entry.city,
            entry.hotel_name,
            float(entry.hotel_amount),
            entry.benefit_type.value,
            float(entry.benefit_amount),
            rdv.status.value,
        ]
        for entry in rdv.entries
    ]


EXPORT_HEADERS = [
    "Protocolo",
    "Colaborador",
    "Função",
    "Data",
    "Cidade",
    "Hotel",
    "Valor hotel",
    "Tipo",
    "Valor benefício",
    "Status",
]


def rdv_to_csv(rdv: RdvSubmission) -> bytes:
    output = StringIO(newline="")
    writer = csv.writer(output, delimiter=";")
    writer.writerow(EXPORT_HEADERS)
    writer.writerows(_entry_rows(rdv))
    return ("\ufeff" + output.getvalue()).encode("utf-8")


def rdvs_to_xlsx(rdvs: Iterable[RdvSubmission]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RDVs"
    sheet.append(EXPORT_HEADERS)
    header_fill = PatternFill("solid", fgColor="B5122B")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for rdv in rdvs:
        for row in _entry_rows(rdv):
            sheet.append(row)
    for row in sheet.iter_rows(min_row=2):
        row[6].number_format = "R$ #,##0.00"
        row[8].number_format = "R$ #,##0.00"
    widths = [16, 34, 18, 14, 24, 28, 16, 16, 18, 14]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def rdv_to_xlsx(rdv: RdvSubmission) -> bytes:
    return rdvs_to_xlsx([rdv])


def _draw_logo(canvas: object, x: float, y: float) -> None:
    from reportlab.lib.utils import ImageReader

    logo_path = Path(__file__).resolve().parent / "assets" / "logo_jr.png"
    canvas.drawImage(
        ImageReader(logo_path),
        x,
        y,
        width=22,
        height=22,
        preserveAspectRatio=True,
        mask="auto",
    )


def _prepare_signature_for_export(signature_data: bytes) -> bytes:
    from PIL import Image, ImageFilter, ImageOps

    with Image.open(BytesIO(signature_data)) as source:
        source.thumbnail((900, 300), Image.Resampling.LANCZOS)
        rgba = source.convert("RGBA")
        white_background = Image.new("RGBA", rgba.size, "white")
        flattened = Image.alpha_composite(white_background, rgba).convert("RGB")
        grayscale = ImageOps.grayscale(flattened)
        ink_mask = grayscale.point(lambda pixel: 255 if pixel < 245 else 0)
        ink_mask = ink_mask.filter(ImageFilter.MaxFilter(5))
        prepared = Image.new("RGB", rgba.size, "white")
        prepared.paste(Image.new("RGB", rgba.size, "black"), mask=ink_mask)
        prepared_stream = BytesIO()
        prepared.save(prepared_stream, format="PNG", optimize=True)
        return prepared_stream.getvalue()


def _draw_signature(
    canvas: object,
    signature_data: bytes,
    x: float,
    y: float,
    width: float,
    height: float,
) -> None:
    from reportlab.lib.utils import ImageReader

    signature = ImageReader(BytesIO(_prepare_signature_for_export(signature_data)))
    image_width, image_height = signature.getSize()
    scale = min(width / image_width, height / image_height)
    draw_width = image_width * scale
    draw_height = image_height * scale
    canvas.drawImage(
        signature,
        x + (width - draw_width) / 2,
        y,
        width=draw_width,
        height=draw_height,
        preserveAspectRatio=True,
        mask="auto",
    )


def rdv_to_pdf(rdv: RdvSubmission) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.pdfgen.canvas import Canvas

    buffer = BytesIO()
    page_width, page_height = landscape(A4)
    canvas = Canvas(buffer, pagesize=(page_width, page_height))
    margin = 25
    _draw_logo(canvas, margin, page_height - 42)
    role_title = (
        "MOTORISTA"
        if rdv.employee.role == EmployeeRole.MOTORISTA
        else "AJUDANTE DE MOTORISTA"
    )
    title = f"RELATÓRIO DE DESPESAS DE VIAGEM - RDV - {role_title}"
    canvas.setFillColor(colors.black)
    canvas.setFont("Helvetica-Bold", 17)
    canvas.drawCentredString(page_width / 2, page_height - 42, title)
    canvas.setFont("Helvetica-Bold", 8.5)
    canvas.drawString(margin, page_height - 57, f"NOME: {rdv.employee.name}")
    canvas.setFont("Helvetica", 8.5)
    period_text = f"DATA QUINZENA (INÍCIO E FINAL): {format_date(rdv.period.start_date)} a {format_date(rdv.period.end_date)}"
    canvas.drawString(page_width / 2, page_height - 57, period_text)
    canvas.drawString(
        margin,
        page_height - 73,
        f"HOUVE ADIANTAMENTO DE DIÁRIA?  {'(X) NÃO' if not rdv.advance_received else '( ) NÃO'}  "
        f"{'(X) SIM' if rdv.advance_received else '( ) SIM'}        NO VALOR DE {format_brl(rdv.advance_amount)}",
    )

    table_top = page_height - 80
    table_bottom = 182
    table_width = page_width - 2 * margin
    is_helper = rdv.employee.role == EmployeeRole.AJUDANTE
    fractions = (
        [0.13, 0.22, 0.22, 0.13, 0.15, 0.15] if is_helper else [0.16, 0.32, 0.26, 0.26]
    )
    headers = (
        [
            "DATA",
            "CIDADE",
            "HOTEL",
            "VALOR HOTEL",
            "DIÁRIA EM VIAGEM",
            "TICKET ALIMENTAÇÃO",
        ]
        if is_helper
        else ["DATA", "CIDADE", "DIÁRIA EM VIAGEM", "TICKET ALIMENTAÇÃO"]
    )
    xs = [margin]
    for fraction in fractions:
        xs.append(xs[-1] + table_width * fraction)
    rows = max(len(rdv.entries) + 1, 2)
    row_height = (table_top - table_bottom) / rows
    canvas.setLineWidth(0.6)
    for x in xs:
        canvas.line(x, table_bottom, x, table_top)
    for row in range(rows + 1):
        y = table_top - row * row_height
        canvas.line(margin, y, page_width - margin, y)
    canvas.setFont("Helvetica-Bold", 9)
    for index, header in enumerate(headers):
        canvas.drawString(xs[index] + 3, table_top - row_height + 6, header)
    canvas.setFont("Helvetica", 8.25)
    for row_index, entry in enumerate(rdv.entries, 1):
        y = table_top - (row_index + 1) * row_height + 6
        date_label = "DOMINGO" if entry.date.weekday() == 6 else format_date(entry.date)
        values: list[str]
        daily_amount = (
            format_brl(entry.benefit_amount)
            if entry.benefit_type == BenefitType.DIARIA
            else ""
        )
        ticket_amount = (
            format_brl(entry.benefit_amount)
            if entry.benefit_type == BenefitType.TICKET
            else ""
        )
        if is_helper:
            values = [
                date_label,
                entry.city,
                entry.hotel_name,
                format_brl(entry.hotel_amount) if entry.hotel_amount else "",
                daily_amount,
                ticket_amount,
            ]
        else:
            values = [date_label, entry.city, daily_amount, ticket_amount]
        for col_index, value in enumerate(values):
            max_width = xs[col_index + 1] - xs[col_index] - 6
            text = str(value)
            while text and stringWidth(text, "Helvetica", 8.25) > max_width:
                text = text[:-1]
            if text != str(value):
                text = text[:-3] + "..."
            canvas.drawString(xs[col_index] + 3, y, text)

    totals = calculate_rdv_totals(rdv.entries, rdv.advance_amount)
    y = table_bottom - 19
    canvas.setFont("Helvetica-Bold", 9)
    canvas.drawString(
        margin,
        y,
        f"TOTAL DA QUINZENA EM R$ -----> {format_brl(totals['expense_total'])}",
    )
    canvas.setFont("Helvetica", 8.5)
    location = getattr(rdv, "location", "") or ""
    signed_date = getattr(rdv, "signed_date", None)
    location_date = (
        f"LOCAL/DATA: {location}, {format_long_date(signed_date)}"
        if location and signed_date
        else "LOCAL/DATA: ____________________________________________, ____ de ____________ de ____________"
    )
    canvas.drawString(
        margin,
        y - 23,
        location_date,
    )
    signature_y = 60
    signature_height = 36
    labels = ["ASSINATURA DO COLABORADOR", "ANALISTA DE FROTA", "GESTOR DE FROTA"]
    signatures = [
        getattr(rdv, "signature_data", None),
        getattr(rdv, "analyst_signature_data", None),
        getattr(rdv, "manager_signature_data", None),
    ]
    signature_width = (table_width - 30) / 3
    for index, label in enumerate(labels):
        start = margin + index * (signature_width + 15)
        canvas.drawCentredString(start + signature_width / 2, signature_y + 48, label)
        canvas.line(start, signature_y, start + signature_width, signature_y)
        if signatures[index]:
            _draw_signature(
                canvas,
                signatures[index],
                start + 3,
                signature_y + 2,
                signature_width - 6,
                signature_height,
            )
    observation = (
        "OBSERVAÇÃO: NOS TERMOS DA CONVENÇÃO COLETIVA, A DIÁRIA DE VIAGEM É DESTINADA AO COLABORADOR QUE EXERCE "
        "ATIVIDADE FORA DA BASE. CONSIDERA-SE CADA PERÍODO MODULAR DE 24 HORAS. O RECEBIMENTO DA DIÁRIA EXCLUI O "
        "PAGAMENTO DA AJUDA DE ALIMENTAÇÃO (TICKET)."
    )
    canvas.setFont("Helvetica", 6.8)
    text = canvas.beginText(margin, 27)
    max_chars = 190
    remaining = observation
    while remaining:
        split = min(max_chars, len(remaining))
        if split < len(remaining):
            split = remaining.rfind(" ", 0, split)
        text.textLine(remaining[:split])
        remaining = remaining[split:].lstrip()
    canvas.drawText(text)
    canvas.setTitle(protocol(rdv.id))
    canvas.save()
    return buffer.getvalue()


def pdf_to_png(pdf_data: bytes, scale: float = 3.0) -> bytes:
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(pdf_data)
    try:
        page = document[0]
        try:
            image = page.render(scale=scale).to_pil()
            output = BytesIO()
            image.save(output, format="PNG", optimize=True)
            return output.getvalue()
        finally:
            page.close()
    finally:
        document.close()


def rdv_to_png(rdv: RdvSubmission) -> bytes:
    return pdf_to_png(rdv_to_pdf(rdv))
